import re
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".docx",
    ".xlsx",
    ".xls",
    ".pptx",
    ".ppt",
}

EXTENSION_TO_LOADER = {
    ".pdf": "pdf",
    ".txt": "text",
    ".md": "text",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".pptx": "pptx",
    ".ppt": "pptx",
}


def _load_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n\n".join(parts)


def _load_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _load_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _load_xlsx(file_path: str) -> str:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join(cells).strip()
            if line:
                parts.append(line)
    wb.close()
    return "\n\n".join(parts)


def _load_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text.strip())
    return "\n\n".join(parts)


def load_from_path(file_path: str, file_type: str | None = None) -> str:
    """Extract text from a file given its path. file_type can override extension."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    loader_type = EXTENSION_TO_LOADER.get(ext) if file_type is None else file_type

    if loader_type == "pdf":
        return _load_pdf(str(path))
    if loader_type == "text":
        return _load_text(str(path))
    if loader_type == "docx":
        return _load_docx(str(path))
    if loader_type == "xlsx":
        return _load_xlsx(str(path))
    if loader_type == "pptx":
        return _load_pptx(str(path))

    raise ValueError(f"Unsupported file type: {ext or loader_type}")


def load_from_url(url: str) -> str:
    """Fetch URL and extract text from HTML."""
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; Jupiter/1.0; +https://github.com/jupiter)",
    }
    response = requests.get(url, headers=headers, timeout=30)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")

    for tag in soup(["script", "style"]):
        tag.decompose()

    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines))
    return text
